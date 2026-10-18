#!/usr/bin/env python3
"""
Shishya — turning a model's answer into a real .xlsx, .docx or .pptx.

Split out of make_app.py so the chat, Training Data, Make a Document and the
Windows build all write files through one implementation: a file built from
any of them is identical to one built from the others.

── Why the model never writes the file ─────────────────────────────────────

The model returns *content* as JSON — rows, headings, slides — and Python
writes the file. Asking a model to emit a document format directly produces
something that looks right and is subtly corrupt, and a file that will not
open in front of a judge is worse than no feature at all. Structure is the
part a model is good at; serialisation is the part a library is good at.
"""

import re
import time

import common


# What we ask for back, per kind. Constraining generation to a schema is what
# keeps this from returning prose where a table was wanted.
SCHEMAS = {
    "Spreadsheet": {
        "type": "object",
        "properties": {
            "filename": {"type": "string"},
            "sheets": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "name": {"type": "string"},
                        "columns": {"type": "array", "items": {"type": "string"}},
                        "rows": {"type": "array",
                                 "items": {"type": "array", "items": {"type": "string"}}},
                    },
                    "required": ["name", "columns", "rows"],
                },
            },
        },
        "required": ["filename", "sheets"],
    },
    "Document": {
        "type": "object",
        "properties": {
            "filename": {"type": "string"},
            "title": {"type": "string"},
            "blocks": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "kind": {"type": "string", "enum": ["heading", "paragraph", "bullet"]},
                        "text": {"type": "string"},
                    },
                    "required": ["kind", "text"],
                },
            },
        },
        "required": ["filename", "title", "blocks"],
    },
    "Presentation": {
        "type": "object",
        "properties": {
            "filename": {"type": "string"},
            "title": {"type": "string"},
            "slides": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "heading": {"type": "string"},
                        "bullets": {"type": "array", "items": {"type": "string"}},
                    },
                    "required": ["heading", "bullets"],
                },
            },
        },
        "required": ["filename", "title", "slides"],
    },
}


ASKS = {
    "Spreadsheet": "Design a spreadsheet. Give real column headings and real "
                   "rows of plausible data — never placeholders like 'Row 1'.",
    "Document": "Write the document. Use headings, paragraphs and bullets as "
                "the content actually calls for.",
    "Presentation": "Build the deck. One idea per slide, three to five short "
                    "bullets each — phrases, not sentences.",
}


def read_training(budget=6000):
    out = []
    if not common.TRAINING.exists():
        return ""
    for p in sorted(common.TRAINING.iterdir()):
        if p.is_file():
            text, _ = common.extract(p)
            out.append(f"[{p.name}]\n{text}")
    return "\n\n".join(out)[:budget]


def safe_name(spec, extension):
    stem = re.sub(r"[^A-Za-z0-9 _-]", "", str(spec.get("filename", "document"))).strip()
    stem = (stem or "document")[:60]
    common.WORK.mkdir(parents=True, exist_ok=True)
    path = common.WORK / f"{stem}{extension}"
    if path.exists():
        path = common.WORK / f"{stem}-{time.strftime('%H%M%S')}{extension}"
    return path


def write_xlsx(spec):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in spec["sheets"] or [{"name": "Sheet1", "columns": [], "rows": []}]:
        ws = wb.create_sheet(title=str(sheet.get("name", "Sheet"))[:31])
        columns = sheet.get("columns", [])
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3864")
        for row in sheet.get("rows", []):
            ws.append(row)
        # Width from content: a spreadsheet whose columns all read ###### is
        # the first thing anyone notices.
        for i, name in enumerate(columns, 1):
            longest = max([len(str(name))] +
                          [len(str(r[i - 1])) for r in sheet.get("rows", [])
                           if len(r) >= i], default=10)
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(48, longest + 3)
        ws.freeze_panes = "A2"
    path = safe_name(spec, ".xlsx")
    wb.save(path)
    return path


def write_docx(spec):
    from docx import Document
    doc = Document()
    doc.add_heading(str(spec.get("title", "Document")), level=0)
    for block in spec.get("blocks", []):
        kind, text = block.get("kind"), str(block.get("text", ""))
        if kind == "heading":
            doc.add_heading(text, level=1)
        elif kind == "bullet":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)
    path = safe_name(spec, ".docx")
    doc.save(path)
    return path


def write_pptx(spec):
    from pptx import Presentation
    from pptx.util import Pt
    deck = Presentation()

    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = str(spec.get("title", "Presentation"))
    if len(cover.placeholders) > 1:
        cover.placeholders[1].text = time.strftime("%d %B %Y")

    for slide_spec in spec.get("slides", []):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = str(slide_spec.get("heading", ""))
        body = slide.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate(slide_spec.get("bullets", [])):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = str(bullet)
            para.font.size = Pt(20)
    path = safe_name(spec, ".pptx")
    deck.save(path)
    return path


WRITERS = {"Spreadsheet": write_xlsx, "Document": write_docx, "Presentation": write_pptx}
