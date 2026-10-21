#!/usr/bin/env python3
"""
Sheets — the clone's spreadsheet. Runs INSIDE the sandbox.

    pythonw sheet_app.py [file.xlsx]

On the Mac the clone's spreadsheet was LibreOffice Calc, and Sheets was a
quick viewer beside it. On Windows Sheets is the spreadsheet: LibreOffice is
not part of the install, and a lesson taught against an application must be
performed against the same one. So it has what a lesson needs from Calc:

  - a Name Box, showing the current cell and taking a reference to go to —
    the route the executor uses, exactly as it used Calc's;
  - a grid with A, B, C… across and 1, 2, 3… down, cells addressed the way a
    person says them;
  - typing into the current cell starts editing it, Enter commits and moves
    down, Tab moves right;
  - every cell readable from outside by row and column through UI
    Automation's Grid pattern, which is how a typed value is verified.

openpyxl reads and writes the real file, so what opens here is the file —
not a preview of it.
"""

import csv
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import kit   # noqa: E402
from PySide6.QtCore import Qt   # noqa: E402
from PySide6.QtWidgets import (QAbstractItemDelegate, QAbstractItemView, QHBoxLayout, QLineEdit,   # noqa: E402
                               QTableWidget, QTableWidgetItem, QTabWidget, QVBoxLayout)

import common   # noqa: E402

MIN_ROWS, MIN_COLS = 60, 16
MAX_ROWS, MAX_COLS = 2000, 60
REF = re.compile(r"^\s*([A-Za-z]{1,3})(\d{1,7})\s*$")


def column_letter(i):
    """0 -> A, 25 -> Z, 26 -> AA."""
    name, i = "", i + 1
    while i:
        i, rem = divmod(i - 1, 26)
        name = chr(65 + rem) + name
    return name


def parse_ref(text):
    m = REF.match(text or "")
    if not m:
        return None
    col = 0
    for letter in m.group(1).upper():
        col = col * 26 + (ord(letter) - 64)
    return int(m.group(2)) - 1, col - 1


class Grid(QTableWidget):
    """One worksheet."""

    def __init__(self, name, rows, owner):
        height = min(max(len(rows) + 20, MIN_ROWS), MAX_ROWS)
        width = min(max(max((len(r) for r in rows), default=0) + 4, MIN_COLS), MAX_COLS)
        super().__init__(height, width)
        self.owner = owner
        self.setObjectName("grid")
        self.setAccessibleName(f"{name} grid")
        self.setHorizontalHeaderLabels([column_letter(c) for c in range(width)])
        self.setVerticalHeaderLabels([str(r + 1) for r in range(height)])
        self.setEditTriggers(QAbstractItemView.AnyKeyPressed | QAbstractItemView.DoubleClicked
                             | QAbstractItemView.EditKeyPressed)
        self.setSelectionMode(QAbstractItemView.SingleSelection)
        self.verticalHeader().setDefaultSectionSize(24)
        self.horizontalHeader().setDefaultSectionSize(110)
        self.blockSignals(True)
        for r, row in enumerate(rows[:MAX_ROWS]):
            for c, value in enumerate(row[:MAX_COLS]):
                if value not in (None, ""):
                    self.setItem(r, c, QTableWidgetItem(str(value)))
        self.blockSignals(False)
        self.currentCellChanged.connect(self.moved)
        self.itemChanged.connect(lambda *_: owner.edited())
        self.setCurrentCell(0, 0)

    def current_ref(self):
        """The name the recorder and the Name Box use for where we are."""
        r, c = self.currentRow(), self.currentColumn()
        return f"{column_letter(max(c, 0))}{max(r, 0) + 1}"

    def moved(self, *_):
        self.owner.show_ref(self)
        kit.Recorder.note("object:selection-changed", self)

    def keyPressEvent(self, e):
        # Enter commits and moves down, Tab moves right — as every spreadsheet.
        if e.key() in (Qt.Key_Return, Qt.Key_Enter) and self.state() != QAbstractItemView.EditingState:
            self.setCurrentCell(min(self.currentRow() + 1, self.rowCount() - 1), self.currentColumn())
            return
        super().keyPressEvent(e)

    def closeEditor(self, editor, hint):
        super().closeEditor(editor, hint)
        # The hints live on QAbstractItemDelegate.EndEditHint. Read from
        # QAbstractItemView they raised AttributeError on every committed
        # cell, so Enter after typing never moved down a row.
        #
        # Only Enter (SubmitModelCache) moves down. NoHint is the editor
        # closing because focus went elsewhere — to the cell-contents line,
        # say, which is how a lesson types — and moving then put the value
        # in the row below the one it was meant for.
        if hint == QAbstractItemDelegate.EndEditHint.SubmitModelCache:
            self.setCurrentCell(min(self.currentRow() + 1, self.rowCount() - 1), self.currentColumn())

    def values(self):
        out = []
        last_row = -1
        for r in range(self.rowCount()):
            row = []
            for c in range(self.columnCount()):
                item = self.item(r, c)
                row.append(item.text() if item else "")
            while row and row[-1] == "":
                row.pop()
            out.append(row)
            if row:
                last_row = r
        return out[:last_row + 1]


class SheetApp(kit.Window):
    def __init__(self, path=None):
        super().__init__("Sheets", 1080, 640)
        self.path = None
        self.book = None
        self.dirty = False

        root = QVBoxLayout(self.body)
        root.setContentsMargins(10, 10, 10, 10)

        bar = QHBoxLayout()
        bar.addWidget(kit.button("Open", self.on_open))
        bar.addWidget(kit.button("Save", self.on_save))
        self.title = kit.label("No file open", role="title")
        bar.addWidget(self.title, 1)
        self.note = kit.label("", role="muted", name="status")
        bar.addWidget(self.note)
        root.addLayout(bar)

        # The Name Box and the value line, as in Calc and Excel.
        refs = QHBoxLayout()
        self.name_box = QLineEdit()
        self.name_box.setObjectName("nameBox")
        self.name_box.setAccessibleName("Name Box")
        self.name_box.setFixedWidth(90)
        self.name_box.returnPressed.connect(self.on_goto)
        refs.addWidget(self.name_box)
        self.value = QLineEdit()
        self.value.setObjectName("cellValue")
        self.value.setAccessibleName("Cell contents")
        self.value.returnPressed.connect(self.on_value)
        refs.addWidget(self.value, 1)
        root.addLayout(refs)

        self.tabs = QTabWidget()
        self.tabs.setObjectName("sheets")
        self.tabs.setAccessibleName("Sheets")
        self.tabs.currentChanged.connect(lambda *_: self.grid() and self.show_ref(self.grid()))
        root.addWidget(self.tabs, 1)

        if path:
            self.load(Path(path))
        else:
            self.fill([("Sheet1", [])])
            self.title.setText("New sheet")

    # ── where we are ──
    def grid(self):
        return self.tabs.currentWidget()

    def show_ref(self, grid):
        if grid is not self.grid():
            return
        self.name_box.setText(grid.current_ref())
        item = grid.currentItem()
        self.value.setText(item.text() if item else "")

    def on_goto(self):
        where = parse_ref(self.name_box.text())
        grid = self.grid()
        if not where or grid is None:
            self.say(f"{self.name_box.text()!r} is not a cell")
            return
        r, c = where
        if r >= grid.rowCount():
            grid.setRowCount(min(r + 20, MAX_ROWS))
            grid.setVerticalHeaderLabels([str(i + 1) for i in range(grid.rowCount())])
        if c >= grid.columnCount():
            grid.setColumnCount(min(c + 4, MAX_COLS))
            grid.setHorizontalHeaderLabels([column_letter(i) for i in range(grid.columnCount())])
        grid.setCurrentCell(r, c)
        grid.scrollToItem(grid.item(r, c) or QTableWidgetItem(), QAbstractItemView.PositionAtCenter)
        grid.setFocus()

    def on_value(self):
        grid = self.grid()
        if grid is None:
            return
        grid.setItem(grid.currentRow(), grid.currentColumn(), QTableWidgetItem(self.value.text()))
        grid.setFocus()

    # ── files ──
    def on_open(self):
        picked = kit.open_files(self, "Open a spreadsheet", common.WORK,
                                "Spreadsheets (*.xlsx *.xlsm *.csv)")
        if picked:
            self.load(picked[0])

    def fill(self, sheets):
        self.tabs.clear()
        for name, rows in sheets:
            self.tabs.addTab(Grid(name, rows, self), name)
        if self.grid():
            self.show_ref(self.grid())

    def load(self, path):
        try:
            if path.suffix.lower() == ".csv":
                with open(path, newline="", encoding="utf-8", errors="replace") as f:
                    self.book, sheets = None, [(path.stem, list(csv.reader(f))[:MAX_ROWS])]
            else:
                from openpyxl import load_workbook
                self.book = load_workbook(path)
                sheets = [(ws.title, [list(r) for r in ws.iter_rows(
                    max_row=MAX_ROWS, max_col=MAX_COLS, values_only=True)])
                    for ws in self.book.worksheets]
        except Exception as e:
            return self.say(f"Could not open {path.name}: {e}")
        self.path = path
        self.fill(sheets)
        self.title.setText(path.name)
        self.setWindowTitle(f"Sheets — {path.name}")
        self.say(f"{len(sheets)} sheet{'s' if len(sheets) != 1 else ''}")
        self.dirty = False

    def on_save(self):
        if self.path is None or self.path.suffix.lower() == ".csv":
            common.WORK.mkdir(parents=True, exist_ok=True)
            name = kit.save_file(self, "Save as", common.WORK / "sheet.xlsx",
                                 "Excel workbook (*.xlsx)")
            if not name:
                return
            if name.suffix.lower() != ".xlsx":
                name = name.with_suffix(".xlsx")
            self.path, self.book = name, None
        try:
            from openpyxl import Workbook
            if self.book is None:
                self.book = Workbook()
                self.book.remove(self.book.active)
            for i in range(self.tabs.count()):
                name = self.tabs.tabText(i)
                ws = self.book[name] if name in self.book.sheetnames else self.book.create_sheet(name)
                grid = self.tabs.widget(i)
                for r, row in enumerate(grid.values(), start=1):
                    for c, value in enumerate(row, start=1):
                        ws.cell(row=r, column=c).value = _typed(value)
            self.book.save(self.path)
            self.dirty = False
            self.title.setText(self.path.name)
            self.say(f"Saved {self.path.name}")
        except Exception as e:
            self.say(f"Could not save: {e}")

    def edited(self):
        self.dirty = True
        self.say("edited — press Save to write it back")
        if self.grid():
            self.show_ref(self.grid())

    def say(self, text):
        self.note.setText(text)


def _typed(value):
    """A number typed as text is stored as a number, as a spreadsheet would."""
    if value in ("", None):
        return None
    try:
        return int(value) if re.fullmatch(r"-?\d+", value) else float(value)
    except (TypeError, ValueError):
        return value


if __name__ == "__main__":
    app = kit.application("Sheets")
    sys.exit(kit.run(SheetApp(sys.argv[1] if len(sys.argv) > 1 else None)))
